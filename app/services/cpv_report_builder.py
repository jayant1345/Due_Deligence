"""
Fill the firm's own CPV report templates (single / dual applicant) with case data.

The templates in app/report_templates/ are the firm's Word reports converted once with
Microsoft Word, so fonts, tables, borders, letterhead and page size are exactly the firm's.
This module only writes values into existing cells/paragraphs (keeping each one's own
formatting) and places photographs on the photo pages. Unknown values stay blank — the
template's sample answers ("Hindu", "Married", "15 Years", "Positive", ...) are never kept
as if they were verified facts.
"""
import copy
import os
from typing import Any, Dict, List, Optional

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

TEMPLATE_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "report_templates")
TEMPLATES = {1: "cpv_1_person.docx", 2: "cpv_2_person.docx"}
PHOTO_SECTIONS = {
    "pan": "Pan card Verification Photographs",
    "residence": "Residence premises Visit photographs",
    "office": "Business / Employer Premises visit",
}
# Usable photo area on a template page (US Letter, firm margins, below the section title)
PHOTO_AREA_W_IN = 6.3
PHOTO_AREA_H_IN = 7.3


def _v(value: Any) -> str:
    return "" if value is None else str(value).strip()


# ---------------------------------------------------------------------------
# Low-level writers that preserve the template's formatting
# ---------------------------------------------------------------------------

def _first_run_props(p) -> Optional[Any]:
    if p.runs and p.runs[0]._r.rPr is not None:
        return copy.deepcopy(p.runs[0]._r.rPr)
    ppr = p._p.pPr
    if ppr is not None and ppr.find(qn("w:rPr")) is not None:
        rpr = copy.deepcopy(ppr.find(qn("w:rPr")))
        rpr.tag = qn("w:rPr")
        return rpr
    return None


def _set_par_text(p, text: str) -> None:
    rpr = _first_run_props(p)
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
        # The template highlights its sample text as "to be edited"; a filled value is final
        rp = p.runs[0]._r.rPr
        if rp is not None:
            for hl in rp.findall(qn("w:highlight")):
                rp.remove(hl)
    else:
        run = p.add_run(text)
        if rpr is not None:
            run._r.insert(0, rpr)


def _set_cell(cell, value: Any) -> None:
    pars = cell.paragraphs
    _set_par_text(pars[0], _v(value))
    for p in pars[1:]:
        _set_par_text(p, "")


def _ucells(table, row: int) -> list:
    """Unique (un-merged) cells of a row, left to right."""
    out, prev = [], None
    for c in table.rows[row].cells:
        if c._tc is prev:
            continue
        prev = c._tc
        out.append(c)
    return out


def _set(table, row: int, col: int, value: Any) -> None:
    _set_cell(_ucells(table, row)[col], value)


def _all_paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _replace_token(doc, token: str, text: Optional[str], drop_if_empty: bool = False) -> None:
    for p in list(_all_paragraphs(doc)):
        if token in p.text:
            if drop_if_empty and not _v(text):
                p._p.getparent().remove(p._p)
            else:
                _set_par_text(p, p.text.replace(token, _v(text)))


def _set_numpages_field(cell) -> None:
    """Total pages as a Word NUMPAGES field (counted by Word itself), keeping the cell's font."""
    p = cell.paragraphs[0]
    rpr = _first_run_props(p)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)

    def _run(child):
        r = OxmlElement("w:r")
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        r.append(child)
        p._p.append(r)

    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); _run(begin)
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = r" NUMPAGES \* MERGEFORMAT "; _run(instr)
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate"); _run(sep)
    t = OxmlElement("w:t"); t.text = cell.text or "1"; _run(t)
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end"); _run(end)


def _page_break_run(p) -> None:
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r = OxmlElement("w:r")
    r.append(br)
    p._p.append(r)


def _is_filler(el) -> bool:
    if el is None or el.tag != qn("w:p"):
        return False
    xml = el.xml
    if "<w:drawing" in xml or "w:txbxContent" in xml or 'w:type="page"' in xml or "w:sectPr" in xml:
        return False
    return not "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def _start_new_page(p) -> None:
    """
    The template pushes each section onto a new page with runs of empty lines sized for its
    sample text; with real (shorter/longer) content those drift. Replace them with an explicit
    page break before the section heading.
    """
    prev = p._p.getprevious()
    while _is_filler(prev):
        nxt = prev.getprevious()
        prev.getparent().remove(prev)
        prev = nxt
    p.paragraph_format.page_break_before = True


# ---------------------------------------------------------------------------
# Photo pages
# ---------------------------------------------------------------------------

def _fit(path: str, max_w: float, max_h: float):
    with Image.open(path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    return Inches(w * scale), Inches(h * scale)


def _deep_text(p) -> str:
    # Includes text inside text boxes anchored in the paragraph (python-docx's .text does not)
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def _fill_photo_section(doc, title: str, photos: List[str], next_titles: List[str], is_last: bool) -> None:
    pars = doc.paragraphs
    start = next((i for i, p in enumerate(pars) if title in _deep_text(p)), None)
    if start is None:
        return
    end = len(pars)
    for i in range(start + 1, len(pars)):
        if any(t in _deep_text(pars[i]) for t in next_titles):
            end = i
            break
    body = [pars[i] for i in range(start + 1, end)
            if not _deep_text(pars[i]).strip() and "<w:drawing" not in pars[i]._p.xml]
    if not body:
        return
    photos = [p for p in photos if p and os.path.exists(p)]
    first = body[0]
    for p in body[1:]:
        p._p.getparent().remove(p._p)
    if not photos:
        _set_par_text(first, "No photographs on record for this section.")
        if is_last:
            _page_break_run(first)
        return

    n = len(photos)
    cols = 1 if n <= 2 else 2
    rows = (n + cols - 1) // cols
    box_w = PHOTO_AREA_W_IN / cols - (0.1 if cols > 1 else 0)
    box_h = PHOTO_AREA_H_IN / rows - 0.15
    first.alignment = 1  # centre
    for k, path in enumerate(photos):
        w, h = _fit(path, box_w, box_h)
        if k and cols == 1:
            first.add_run().add_break()
        elif k and k % cols == 0:
            first.add_run().add_break()
        elif k:
            first.add_run("  ")
        first.add_run().add_picture(path, width=w, height=h)
    if is_last:
        _page_break_run(first)


# ---------------------------------------------------------------------------
# Section fillers
# ---------------------------------------------------------------------------

def _fill_identity(table, person: Dict[str, Any]) -> None:
    _set(table, 0, 1, person.get("display_name"))
    _set(table, 1, 1, person.get("contact"))
    row2 = _ucells(table, 2)
    _set_cell(row2[1], person.get("dob"))
    _set_cell(row2[3], person.get("age"))
    _set_cell(row2[5], person.get("gender"))
    _set(table, 3, 1, person.get("marital_status"))
    _set(table, 4, 1, person.get("community"))
    _set_cell(_ucells(table, 5)[-1], person.get("pan"))
    _set_cell(_ucells(table, 6)[-1], person.get("aadhaar"))
    _set(table, 7, 1, person.get("qualification"))
    _set(table, 8, 1, person.get("occupation"))


def _fill_family(table, family: List[Dict[str, Any]]) -> None:
    for r in range(1, len(table.rows)):
        m = family[r - 1] if r - 1 < len(family) else {}
        cells = _ucells(table, r)
        _set_cell(cells[-4], m.get("name"))
        _set_cell(cells[-3], m.get("relation"))
        _set_cell(cells[-2], m.get("age"))
        _set_cell(cells[-1], m.get("earning"))


def _fill_residence(t_premises, t_details, res: Dict[str, Any]) -> None:
    _set(t_premises, 0, 1, res.get("verified_address"))
    _set(t_premises, 1, 1, res.get("aadhaar_address"))
    _set(t_premises, 2, 1, res.get("landmark"))
    refs = res.get("references") or []
    for i, r in enumerate((4, 5)):
        name, rel = (refs[i] if i < len(refs) else ("", ""))
        cells = _ucells(t_premises, r)
        _set_cell(cells[-2], name)
        _set_cell(cells[-1], rel)
    row6 = _ucells(t_premises, 6)
    _set_cell(row6[1], res.get("building_type"))
    _set_cell(row6[3], res.get("floor"))
    _set_cell(_ucells(t_premises, 7)[-1], res.get("stay_period"))
    _set(t_premises, 8, 1, res.get("ownership"))
    _set(t_premises, 9, 1, res.get("owner_name"))
    _set(t_premises, 10, 1, res.get("owner_relation"))

    for r, key in enumerate(("nameboard", "nature", "appearance", "locality", "surroundings",
                             "assets", "evidence", "assessment")):
        _set_cell(_ucells(t_details, r)[-1], res.get(key))
    met = res.get("met") or {}
    _set_cell(_ucells(t_details, 8)[-1], met.get("name"))
    _set_cell(_ucells(t_details, 9)[-1], met.get("contact"))
    _set_cell(_ucells(t_details, 10)[-1], met.get("relation"))


def _fill_office(table, person: Dict[str, Any]) -> None:
    o = person.get("office") or {}
    _set(table, 0, 1, person.get("display_name"))
    _set(table, 1, 1, o.get("firm_name"))
    _set(table, 2, 1, o.get("address"))
    _set(table, 3, 1, o.get("landmark"))
    _set(table, 4, 1, o.get("firm_type"))
    _set(table, 5, 1, o.get("premises_ownership"))
    row6 = _ucells(table, 6)
    _set_cell(row6[1], o.get("premises_owner"))
    _set_cell(row6[3], o.get("owner_relation"))
    row7 = _ucells(table, 7)
    _set_cell(row7[1], o.get("area"))
    _set_cell(row7[3], o.get("period_of_stay"))
    row8 = _ucells(table, 8)
    _set_cell(row8[1], o.get("established_since"))
    _set_cell(row8[3], o.get("experience"))
    row9 = _ucells(table, 9)
    _set_cell(row9[1], o.get("location"))
    _set_cell(row9[3], o.get("employees"))
    _set(table, 10, 1, o.get("nameboard"))
    _set(table, 11, 1, o.get("infrastructure"))
    _set(table, 12, 1, o.get("reg_or_emp_no"))
    _set(table, 13, 1, o.get("position"))
    _set(table, 14, 1, o.get("assets"))
    refs = o.get("references") or []
    for i, r in enumerate((16, 17)):
        name, rel = (refs[i] if i < len(refs) else ("", ""))
        cells = _ucells(table, r)
        _set_cell(cells[-2], name)
        _set_cell(cells[-1], rel)
    _set(table, 18, 1, o.get("assessment"))
    met = o.get("met") or {}
    _set_cell(_ucells(table, 19)[-1], met.get("name"))
    _set_cell(_ucells(table, 20)[-1], met.get("contact"))
    _set_cell(_ucells(table, 21)[-1], met.get("relation"))


def _fill_itr(table, rows: List[Dict[str, Any]]) -> None:
    for i, r in enumerate((1, 3, 5)):
        item = rows[i] if i < len(rows) else {}
        cells = _ucells(table, r)
        _set_cell(cells[0], item.get("ay"))
        _set_cell(cells[1], item.get("ack_no"))
        _set_cell(cells[2], item.get("filed_on"))
        _set_cell(cells[3], ("Rs. " + item["total_income"]) if item.get("total_income") else "Rs.")
        _set_cell(_ucells(table, r + 1)[-1], item.get("status"))


def _fill_bank(table, bank: Dict[str, Any]) -> None:
    for r, key in enumerate(("bank_name", "account_number", "branch", "account_status", "account_type")):
        _set_cell(_ucells(table, r)[-1], bank.get(key))


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def build_cpv_report(data: Dict[str, Any], output_path: str) -> Dict[str, Any]:
    applicant = data["applicant"]
    co = data.get("co_applicant")
    persons = 2 if co else 1
    doc = docx.Document(os.path.join(TEMPLATE_DIR, TEMPLATES[persons]))
    T = doc.tables

    # Cover page
    _replace_token(doc, "[[LOAN_TYPE]]", data.get("loan_type"))
    _set(T[0], 0, 1, data.get("bank_branch"))
    _set(T[0], 1, 1, data.get("report_date"))
    _set_cell(_ucells(T[0], 2)[-1], applicant.get("display_name"))
    row = 3
    if co:
        _set_cell(_ucells(T[0], 3)[-1], co.get("display_name"))
        row = 4
    verifiers = (data.get("verifier_names") or []) + ["", ""]
    _set_cell(_ucells(T[0], row)[-1], verifiers[0])
    _set_cell(_ucells(T[0], row + 1)[-1], verifiers[1])
    _set_cell(_ucells(T[0], row + 2)[-1], data.get("file_allotted_by"))
    _set_numpages_field(_ucells(T[1], 6)[-1])
    _set_cell(_ucells(T[2], 1)[-1], data.get("overall_status"))

    # Summary narrative
    narrative = data.get("narrative") or {}
    for key, count in (("RES", 4), ("WORK", 3), ("ITR", 2), ("KYC", persons)):
        lines = narrative.get(key.lower()) or []
        for i in range(count):
            _replace_token(doc, f"[[{key}_{i + 1}]]", lines[i] if i < len(lines) else None, drop_if_empty=True)

    # Annexure A / B / C / D
    if co:
        _fill_identity(T[3], applicant)
        _fill_identity(T[4], co)
        family_t, prem_t, det_t = T[5], T[6], T[7]
        _fill_office(T[8], applicant)
        _fill_office(T[9], co)
        _fill_itr(T[10], applicant.get("itr") or [])
        _fill_itr(T[11], co.get("itr") or [])
        _fill_bank(T[12], applicant.get("bank") or {})
        _fill_bank(T[13], co.get("bank") or {})
    else:
        _fill_identity(T[3], applicant)
        family_t, prem_t, det_t = T[4], T[5], T[6]
        _fill_office(T[7], applicant)
        _fill_itr(T[8], applicant.get("itr") or [])
        _fill_bank(T[9], applicant.get("bank") or {})
    _fill_family(family_t, data.get("family") or [])
    _fill_residence(prem_t, det_t, data.get("residence") or {})
    _replace_token(doc, "[[ITR_NAME_1]]", applicant.get("display_name"))
    _replace_token(doc, "[[BANK_NAME_1]]", applicant.get("display_name"))
    if co:
        _replace_token(doc, "[[ITR_NAME_2]]", co.get("display_name"))
        _replace_token(doc, "[[BANK_NAME_2]]", co.get("display_name"))

    # Photo pages
    photos = data.get("photos") or {}
    titles = list(PHOTO_SECTIONS.values())
    for i, (key, title) in enumerate(PHOTO_SECTIONS.items()):
        _fill_photo_section(doc, title, photos.get(key) or [], titles[i + 1:] + ["FINAL CERTIFICATE"], i == len(titles) - 1)

    # Section page starts, per the template's page plan
    starts = ["__________Annexure ‘A’", "__________Annexure ‘B’", "__________Annexure ‘C’"]
    if co:
        starts.append("__________Annexure ‘D’")
    pars = doc.paragraphs
    for marker in starts:
        p = next((x for x in pars if x.text.strip().startswith(marker)), None)
        if p is not None:
            _start_new_page(p)
    if co:  # the co-applicant's office verification has its own page
        offices = [x for x in pars if x.text.strip().startswith("Office Verification (Business/ Employer)")]
        if len(offices) > 1:
            _start_new_page(offices[1])
    for title in PHOTO_SECTIONS.values():
        p = next((x for x in doc.paragraphs if title in _deep_text(x)), None)
        if p is not None:
            _start_new_page(p)

    # Final certificate
    _replace_token(doc, "[[UDIN]]", data.get("udin"))
    _replace_token(doc, "[[CERT_DATE]]", data.get("certificate_date"))
    _replace_token(doc, "[[CERT_PLACE]]", data.get("place") or "Ahmedabad")
    sig = data.get("signature_path")
    if sig and os.path.exists(sig):
        for p in _all_paragraphs(doc):
            if "[[SIGNATURE]]" in p.text:
                _set_par_text(p, p.text.replace("[[SIGNATURE]]", ""))
                p.add_run().add_picture(sig, width=Inches(1.6))
                break
    else:
        _replace_token(doc, "[[SIGNATURE]]", "")

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    return {"path": output_path, "template": TEMPLATES[persons], "persons": persons}
